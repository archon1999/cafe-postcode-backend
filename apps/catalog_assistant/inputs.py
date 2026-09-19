"""Bounded clipboard/file normalization. Workbook formulas are never executed."""
import base64
import io
import zipfile
from pathlib import Path

from PIL import Image, UnidentifiedImageError
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT = 40000


def normalize_input(text='', files=()):
    text = str(text or '').strip()
    if len(text) > MAX_TEXT or len(files) > 5:
        raise ValidationError('Bir safarda 40 000 belgigacha matn va 5 tagacha fayl yuboring.')
    content = []
    for upload in files:
        if upload.size > MAX_FILE_BYTES:
            raise ValidationError('Har bir fayl 10 MB dan kichik bo‘lishi kerak.')
        raw = upload.read(MAX_FILE_BYTES + 1)
        if len(raw) > MAX_FILE_BYTES:
            raise ValidationError('Fayl juda katta.')
        suffix = Path(upload.name).suffix.lower()
        if suffix in ('.xlsx', '.csv', '.tsv', '.txt'):
            try:
                if suffix == '.xlsx':
                    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                        if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                            raise ValidationError('Excel faylining ochilgan hajmi juda katta.')
                    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                    try:
                        lines = []
                        if len(book.worksheets) > 5:
                            raise ValidationError('Excel faylida 5 tagacha jadval yuboring.')
                        for sheet in book.worksheets:
                            if (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 20:
                                raise ValidationError('Har bir jadval 500 qator va 20 ustundan oshmasin.')
                            lines.append(f'Jadval: {sheet.title}')
                            for row in sheet.iter_rows(max_row=500, max_col=20, values_only=True):
                                if any(value is not None for value in row):
                                    lines.append('\t'.join(str(value if value is not None else '') for value in row))
                        extracted = '\n'.join(lines)
                    finally:
                        book.close()
                else:
                    extracted = raw.decode('utf-8-sig')
                text += '\n' + extracted
            except (ValueError, OSError, KeyError, zipfile.BadZipFile) as exc:
                raise ValidationError('Jadvalni o‘qib bo‘lmadi. XLSX yoki UTF-8 CSV yuboring.') from exc
        else:
            try:
                with Image.open(io.BytesIO(raw)) as img:
                    if img.width * img.height > 25_000_000:
                        raise ValidationError('Rasm o‘lchami juda katta.')
                    img.load()
                    img.thumbnail((2200, 2200))
                    output = io.BytesIO()
                    img.convert('RGB').save(output, format='JPEG', quality=88)
                encoded = base64.b64encode(output.getvalue()).decode()
                content.append({'type': 'input_image', 'image_url': f'data:image/jpeg;base64,{encoded}'})
            except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
                raise ValidationError('JPG, PNG, WEBP rasm yoki XLSX/CSV/TXT fayl yuboring.') from exc
    if len(text) > MAX_TEXT:
        raise ValidationError('Jadval juda katta. Uni kichikroq qismlarga bo‘lib yuboring.')
    if text:
        content.insert(0, {'type': 'input_text', 'text': text})
    if not content:
        raise ValidationError('Matn, Excel qatorlari yoki rasmni joylang.')
    return content
